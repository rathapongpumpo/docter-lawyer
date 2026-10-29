"""Text extraction and clause segmentation for uploaded contracts (M12, FR-CR-002).

Pure functions: bytes in, structured text out — no I/O, no model. The result is stored encrypted (never in kb.chunks)
and is the only input a later review step sees, so the clause references produced here ("ข้อ 7", page numbers) are
what findings cite. Scanned files (no text layer) are detected and reported, not guessed at: OCR is out of scope
for P1 (FR-CR-013).
"""

import io
import re
from dataclasses import dataclass, field

# Below this many characters per page a PDF is treated as a scan (no usable text layer).
MIN_CHARS_PER_PAGE = 120
_THAI_DIGITS = str.maketrans("๐๑๒๓๔๕๖๗๘๙", "0123456789")


class UnsupportedFile(Exception):
    """The file cannot be read (corrupt, encrypted, wrong type)."""


@dataclass
class Page:
    n: int
    text: str


@dataclass
class Clause:
    ref: str                 # "ข้อ 7", "Article 2" or "ส่วนต้น" (text before the first clause)
    title: str
    page: int | None
    text: str


@dataclass
class Extracted:
    pages: list[Page]
    clauses: list[Clause] = field(default_factory=list)
    is_scan: bool = False
    parties: list[str] = field(default_factory=list)   # company names near the top — a hint, never verified
    page_count: int | None = None

    def as_json(self) -> dict:
        return {"pages": [p.__dict__ for p in self.pages], "clauses": [c.__dict__ for c in self.clauses],
                "parties": self.parties, "page_count": self.page_count}


# Many Thai fonts (TH Sarabun New among them) draw tone marks and vowels shifted to sit above tall consonants and
# store those glyphs in the Private Use Area (U+F700…). A PDF made with such a font extracts as "จ\uf70bาง" instead
# of "จ้าง", which would break clause detection and every later comparison. The table comes from the glyph names
# of TH Sarabun New (uni0E49.alt2 → U+0E49, …); variants of other fonts not listed here are dropped, not kept.
_PUA_THAI = {0xF700: 0x0E10, 0xF701: 0x0E34, 0xF702: 0x0E35, 0xF703: 0x0E36, 0xF704: 0x0E37,
             0xF70F: 0x0E0D, 0xF710: 0x0E31, 0xF711: 0x0E4D, 0xF712: 0x0E47,
             0xF718: 0x0E38, 0xF719: 0x0E39, 0xF71A: 0x0E3A, 0xF71B: 0x0E0E, 0xF71C: 0x0E0F, 0xF71D: 0x0E2C}
for _base in (0xF705, 0xF70A, 0xF713):          # mai ek, mai tho, mai tri, mai chattawa, thanthakhat in three positions
    for _i in range(5):
        _PUA_THAI[_base + _i] = 0x0E48 + _i
_PUA_TABLE = {**_PUA_THAI, **{cp: None for cp in range(0xE000, 0xF900) if cp not in _PUA_THAI}}


_COMBINING = "\u0E31\u0E34-\u0E3A\u0E47-\u0E4E"


def clean(text: str) -> str:
    """Undo PDF extraction artefacts: shifted Thai glyphs (PUA), nikhahit + sara aa (sara am), zero-width marks."""
    text = text.translate(_PUA_TABLE)
    text = re.sub("ํ([\u0E48-\u0E4B])", r"\1ํ", text)          # nikhahit placed before its tone mark: "ซํ้า" → "ซ้ำ"
    text = text.replace("ํา", "ำ").replace("\u200b", "").replace("\ufeff", "")
    text = re.sub(f" +([{_COMBINING}])", r"\1", text)  # a mark drawn as its own glyph leaves a gap: "เป ็น"
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    from app.contracts.pdf_marks import neutralise_unnamed_marks, resolve_marks

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnsupportedFile("encrypted pdf")
        placeholders = neutralise_unnamed_marks(reader)
        raw = [p.extract_text() or "" for p in reader.pages]
        if placeholders:  # decided over the whole document: more occurrences make the choice reliable
            raw = resolve_marks("\f".join(raw), placeholders).split("\f")
        pages = [Page(i + 1, clean(t)) for i, t in enumerate(raw)]
    except PyPdfError as exc:
        raise UnsupportedFile("unreadable pdf") from exc
    total = sum(len(p.text) for p in pages)
    scan = not pages or total / len(pages) < MIN_CHARS_PER_PAGE
    return Extracted(pages, is_scan=scan, page_count=len(pages))


def extract_docx(data: bytes) -> Extracted:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        doc = Document(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise UnsupportedFile("unreadable docx") from exc
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # tables carry prices and schedules in many contracts
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    text = clean("\n".join(lines))
    # a .docx has no fixed pages: everything is page 1 and clauses are located by number only
    return Extracted([Page(1, text)], is_scan=len(text) < 80, page_count=None)


# ------------------------------------------------------------------ clause segmentation
_THAI_CLAUSE = re.compile(r"^\s*(ข้อ\s*([0-9๐-๙]{1,3}))\s*[.):：]?\s*(.*)$")
_NUMBERED = re.compile(r"^\s*([0-9๐-๙]{1,2})\s*[.)]\s+(\S.*)$")
_EN_CLAUSE = re.compile(r"^\s*((?:Article|Section|Clause)\s+([0-9]{1,3}))\s*[.):：-]?\s*(.*)$", re.IGNORECASE)


def _norm_ref(label: str) -> str:
    return re.sub(r"\s+", " ", label.translate(_THAI_DIGITS)).strip()


def segment(pages: list[Page]) -> list[Clause]:
    """Top-level clauses. Sub-clauses (1.1, (a) …) stay inside their parent so a finding can quote the whole clause.
    Thai "ข้อ N" / English "Article|Section|Clause N" win; only when none is found do plain "1." numbers count."""
    lines = [(p.n, ln) for p in pages for ln in p.text.split("\n")]
    for pattern, kind in ((_THAI_CLAUSE, "clause"), (_EN_CLAUSE, "clause"), (_NUMBERED, "numbered")):
        hits = [i for i, (_, ln) in enumerate(lines) if pattern.match(ln)]
        if kind == "numbered" and len(hits) < 3:
            whole = "\n".join(ln for _, ln in lines).strip()
            return [Clause("เนื้อหาทั้งฉบับ", "", pages[0].n, whole)] if lines else []
        if len(hits) < 2 and kind == "clause":
            continue
        return _cut(lines, hits, pattern, kind)
    return []


def _cut(lines: list[tuple[int, str]], hits: list[int], pattern: re.Pattern[str], kind: str) -> list[Clause]:
    out: list[Clause] = []
    first = hits[0]
    if first > 0:
        pre = "\n".join(ln for _, ln in lines[:first]).strip()
        if pre:
            out.append(Clause("ส่วนต้น", "", lines[0][0], pre))
    for k, start in enumerate(hits):
        end = hits[k + 1] if k + 1 < len(hits) else len(lines)
        m = pattern.match(lines[start][1])
        assert m is not None
        if kind == "numbered":
            ref, rest = f"ข้อ {m.group(1).translate(_THAI_DIGITS)}", m.group(2)
        else:
            ref, rest = _norm_ref(m.group(1)), m.group(3)
            ref = re.sub(r"^(Article|Section|Clause)", lambda x: x.group(1).capitalize(), ref, flags=re.IGNORECASE)
        body = "\n".join([rest, *[ln for _, ln in lines[start + 1:end]]]).strip()
        title = rest.strip()[:120] if len(rest.strip()) <= 120 else ""
        out.append(Clause(ref, title, lines[start][0], body))
    return out


_COMPANY = re.compile(r"((?:บริษัท|หจก\.|ห้างหุ้นส่วนจำกัด)\s*[^\n,;()]{2,70}?(?:จำกัด(?:\s*\(มหาชน\))?|จํากัด))")


def find_parties(pages: list[Page]) -> list[str]:
    head = "\n".join(p.text for p in pages[:2])[:4000]
    names = list(dict.fromkeys(re.sub(r"\s+", " ", m.group(1)).strip() for m in _COMPANY.finditer(head)))
    return names[:2]


def extract(data: bytes, mime: str) -> Extracted:
    ex = extract_pdf(data) if mime == "application/pdf" else extract_docx(data)
    if not ex.is_scan:
        ex.clauses = segment(ex.pages)
        ex.parties = find_parties(ex.pages)
    return ex
