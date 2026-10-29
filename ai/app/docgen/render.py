"""DOCX / PDF letters for the template wizard (docs/03-technical-spec.md §7.12, FR-TPL-006/008).

Stateless: the caller sends the finished lines, this module lays them out. Nothing is stored or logged.
PDF: fpdf2 with HarfBuzz shaping (Thai tone marks and vowels position correctly) and PyThaiNLP word breaks, because
Thai has no spaces to wrap on. DOCX: python-docx, TH Sarabun New 16 pt (Word substitutes if the font is not installed).
"""

import io
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from fpdf import FPDF

FONT_DIR = Path(__file__).parent / "fonts"
FONT_NAME = "TH Sarabun New"
BODY_PT, FOOTER_PT = 16, 12
MARGIN_MM = 25


def render_docx(title: str, lines: list[str], footer: str) -> bytes:
    doc = Document()
    doc.core_properties.title = title[:200]
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(MARGIN_MM)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(BODY_PT)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT_NAME)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    for line in lines:
        p = doc.add_paragraph()
        if line:
            p.add_run(line)

    if footer:
        fp = sec.footer.paragraphs[0]
        run = fp.add_run(footer)
        run.font.size = Pt(FOOTER_PT)
        run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        run.font.name = FONT_NAME

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


class _Letter(FPDF):
    footer_text = ""

    def footer(self) -> None:  # called by fpdf2 at every page end
        if not self.footer_text:
            return
        self.set_y(-18)
        self.set_font(FONT_NAME, size=FOOTER_PT)
        self.set_text_color(89, 89, 89)
        for line in _wrap(self, self.footer_text, self.w - 2 * MARGIN_MM):
            self.cell(0, 5, line, new_x="LMARGIN", new_y="NEXT")
        self.set_text_color(0, 0, 0)


def _wrap(pdf: FPDF, text: str, width_mm: float) -> list[str]:
    """Greedy line breaking on PyThaiNLP words (fpdf2 alone would only break at spaces).

    Word widths are measured once and added up: re-measuring the growing line for every word made a page cost seconds.
    Words never merge glyphs across a break, so the sum equals the width of the joined text."""
    from pythainlp.tokenize import word_tokenize

    widths: dict[str, float] = {}

    def width(t: str) -> float:
        if t not in widths:
            widths[t] = pdf.get_string_width(t)
        return widths[t]

    lines: list[str] = []
    for raw in text.split("\n"):
        if not raw:
            lines.append("")
            continue
        current, used = "", 0.0
        for tok in word_tokenize(raw, engine="newmm", keep_whitespace=True):
            if current and used + width(tok) > width_mm:
                lines.append(current.rstrip())
                tok = tok.lstrip()
                current, used = "", 0.0
            # a single token wider than the line (a long URL) is split by characters
            while width(tok) > width_mm - used and len(tok) > 1:
                cut = len(tok)
                while cut > 1 and width(tok[:cut]) > width_mm - used:
                    cut -= 1
                lines.append((current + tok[:cut]).rstrip())
                current, used, tok = "", 0.0, tok[cut:]
            current += tok
            used += width(tok)
        lines.append(current.rstrip())
    return lines


def render_pdf(title: str, lines: list[str], footer: str) -> bytes:
    pdf = _Letter(orientation="P", unit="mm", format="A4")
    pdf.footer_text = footer
    pdf.set_margins(MARGIN_MM, MARGIN_MM, MARGIN_MM)
    pdf.set_auto_page_break(auto=True, margin=24)
    pdf.set_title(title[:200])
    pdf.add_font(FONT_NAME, "", str(FONT_DIR / "THSarabunNew-Regular.ttf"))
    pdf.add_font(FONT_NAME, "B", str(FONT_DIR / "THSarabunNew-Bold.ttf"))
    pdf.set_text_shaping(True)
    pdf.add_page()
    pdf.set_font(FONT_NAME, size=BODY_PT)
    line_h = BODY_PT * 0.3528 * 1.05  # pt → mm with a slightly open leading (TH Sarabun's ascent is tall)
    width = pdf.w - 2 * MARGIN_MM
    for line in lines:
        for row in _wrap(pdf, line, width):
            pdf.cell(0, line_h, row, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())
