"""Template wizard output (FR-TPL-006/008) and the knowledge-search endpoint's auth/validation (FR-KB-009)."""

import io
from types import SimpleNamespace

from docx import Document
from pypdf import PdfReader

from app.docgen.render import render_docx, render_pdf
from app.rag.retrieve import Hit, Retrieval
from tests.conftest import TOKEN

LINES = ["บริษัท ตัวอย่าง จำกัด", "วันที่ 24 กันยายน 2569", "", "เรื่อง  บอกกล่าวให้ปฏิบัติตามสัญญา", "ขอแสดงความนับถือ"]
FOOTER = "แบบตัวอย่างเพื่อการศึกษา ควรให้ผู้เชี่ยวชาญตรวจก่อนใช้กับธุรกรรมมูลค่าสูง"


def test_docx_keeps_every_line_and_the_disclaimer_footer():
    doc = Document(io.BytesIO(render_docx("t", LINES, FOOTER)))
    assert [p.text for p in doc.paragraphs] == LINES
    assert doc.sections[0].footer.paragraphs[0].text == FOOTER
    assert doc.styles["Normal"].font.name == "TH Sarabun New"


def test_pdf_is_valid_and_long_thai_text_wraps_inside_the_page():
    long_line = "ตามสัญญาที่อ้างถึง ท่านมีหน้าที่ส่งมอบสินค้าให้ครบถ้วน" * 12  # no break opportunity except word breaks
    data = render_pdf("t", [long_line, *LINES] * 8, FOOTER)
    reader = PdfReader(io.BytesIO(data))
    assert data.startswith(b"%PDF") and len(reader.pages) >= 2
    assert reader.metadata is not None and reader.metadata.title == "t"


async def test_render_endpoint(client):
    body = {"title": "t", "lines": LINES, "footer": FOOTER, "format": "docx"}
    assert (await client.post("/v1/documents/render", json=body)).status_code == 401
    ok = await client.post("/v1/documents/render", json=body, headers={"X-Service-Token": TOKEN})
    assert ok.status_code == 200 and ok.headers["cache-control"] == "no-store"
    assert ok.headers["content-type"].startswith("application/vnd.openxmlformats")
    bad = await client.post("/v1/documents/render", json={**body, "format": "exe"}, headers={"X-Service-Token": TOKEN})
    assert bad.status_code == 422
    huge = await client.post("/v1/documents/render", json={**body, "lines": ["x"] * 401},
                             headers={"X-Service-Token": TOKEN})
    assert huge.status_code == 422


class _FakeRetriever:
    def __init__(self, hits: list[Hit], requested: list[tuple[str | None, str]]) -> None:
        self.hits, self.requested = hits, requested

    async def search(self, query, *, as_of, include_drafts, scope=()):
        return Retrieval(self.hits, self.requested, False, 1, {})


def _hit(i, source_type, url, dense=0.7, lexical=None, pinned=False, **meta):
    return Hit(i, source_type, i, 1, None, f"t{i}", "เนื้อหา " * 60, {"url": url, "status": "draft", **meta},
               dense=dense, lexical=lexical, pinned=pinned)


async def test_search_filters_noise_dedupes_by_document_and_reports_missing_sections(client):
    from app.main import app

    hits = [_hit(1, "law_section", "/law/ccc/420", pinned=True, law_code="ccc", section_no="420"),
            _hit(2, "article", "/library/articles/a", doc_title="A", heading="H"),
            _hit(3, "article", "/library/articles/a", doc_title="A", heading="H2"),  # same document
            _hit(4, "guide", "/guides/g", dense=0.2),                                 # semantic noise
            _hit(5, "guide", "/guides/g2", dense=0.2, lexical=0.3)]                   # word match keeps it
    app.state.answerer = SimpleNamespace(retriever=_FakeRetriever(hits, [("ccc", "420"), ("ccc", "1336/1")]))
    try:
        r = await client.post("/v1/kb/search", json={"q": "ปพพ 420 และ 1336/1"}, headers={"X-Service-Token": TOKEN})
        body = r.json()
        assert r.status_code == 200
        assert [x["url"] for x in body["results"]] == ["/law/ccc/420", "/library/articles/a", "/guides/g2"]
        assert body["results"][0]["exact"] is True and body["results"][1]["title"] == "A"
        assert body["missing_sections"] == ["ccc 1336/1"]
        only = await client.post("/v1/kb/search", json={"q": "x", "kinds": ["guide"]},
                                 headers={"X-Service-Token": TOKEN})
        assert [x["url"] for x in only.json()["results"]] == ["/guides/g2"]
    finally:
        app.state.answerer = None


async def test_metrics_endpoint_is_closed_without_a_token_and_counts_by_route(client, monkeypatch):
    monkeypatch.delenv("METRICS_TOKEN", raising=False)
    assert (await client.get("/metrics")).status_code == 404
    monkeypatch.setenv("METRICS_TOKEN", "scrape-secret")
    assert (await client.get("/metrics")).status_code == 401
    await client.post("/v1/documents/render", json={"title": "t", "lines": ["x"], "footer": "", "format": "docx"},
                      headers={"X-Service-Token": TOKEN})
    body = await client.get("/metrics", headers={"Authorization": "Bearer scrape-secret"})
    assert body.status_code == 200
    assert 'lc_ai_http_requests_total{route="/v1/documents/render",status="200"}' in body.text
